"""
exporter.py — Generate downloadable CSV and XLSX from scrape results.
"""

import io
import csv
from datetime import datetime

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

from scraper import ScrapedResult, ErrorResult


RESULT_COLUMNS = ["#", "URL", "Vendor", "Product", "Edition", "Version", "Licence Metric", "EOS", "EOL"]
ERROR_COLUMNS  = ["#", "URL", "Error Reason"]


# ---------------------------------------------------------------------------
# CSV Export
# ---------------------------------------------------------------------------

def to_csv(results: list[ScrapedResult], errors: list[ErrorResult]) -> bytes:
    """Build a CSV with results and errors in two sections."""
    buf = io.StringIO()
    writer = csv.writer(buf)

    # Results section
    writer.writerow(["=== SCRAPE RESULTS ==="])
    writer.writerow(RESULT_COLUMNS)
    for i, r in enumerate(results, 1):
        writer.writerow([i, r.url, r.vendor, r.product, r.edition, r.version, r.licence_metric, r.eos, r.eol])

    writer.writerow([])

    # Errors section
    writer.writerow(["=== ERROR URLS ==="])
    writer.writerow(ERROR_COLUMNS)
    for i, e in enumerate(errors, 1):
        writer.writerow([i, e.url, e.error])

    return buf.getvalue().encode("utf-8-sig")  # BOM for Excel compatibility


# ---------------------------------------------------------------------------
# XLSX Export — styled workbook
# ---------------------------------------------------------------------------

def _make_border(color="D1D5DB"):
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def _style_header_row(ws, row_num: int, num_cols: int, bg_color: str, font_color: str = "FFFFFF"):
    """Apply header styling to a row."""
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = Font(bold=True, color=font_color, size=11, name="Calibri")
        cell.fill = PatternFill("solid", fgColor=bg_color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _make_border("9CA3AF")


def _style_data_row(ws, row_num: int, num_cols: int, alternate: bool):
    bg = "F9FAFB" if alternate else "FFFFFF"
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = _make_border("E5E7EB")
        cell.font = Font(size=10, name="Calibri")


def _auto_fit_columns(ws, min_width=12, max_width=60):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            try:
                val_len = len(str(cell.value or ""))
                if val_len > max_len:
                    max_len = val_len
            except Exception:
                pass
        adjusted = max(min_width, min(max_width, max_len + 4))
        ws.column_dimensions[col_letter].width = adjusted


def to_xlsx(results: list[ScrapedResult], errors: list[ErrorResult]) -> bytes:
    """Build a styled XLSX with two sheets: Results and Errors."""
    wb = openpyxl.Workbook()

    # -----------------------------------------------------------------------
    # Sheet 1: Results
    # -----------------------------------------------------------------------
    ws_results = wb.active
    ws_results.title = "Scrape Results"
    ws_results.sheet_view.showGridLines = False
    ws_results.row_dimensions[1].height = 20
    ws_results.row_dimensions[2].height = 30

    # Title row
    ws_results.merge_cells(f"A1:{get_column_letter(len(RESULT_COLUMNS))}1")
    title_cell = ws_results.cell(row=1, column=1)
    title_cell.value = f"Web Scrape Results — {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    title_cell.font = Font(bold=True, size=13, color="1E1B4B", name="Calibri")
    title_cell.fill = PatternFill("solid", fgColor="EDE9FE")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Header row
    for col_idx, col_name in enumerate(RESULT_COLUMNS, 1):
        ws_results.cell(row=2, column=col_idx, value=col_name)
    _style_header_row(ws_results, 2, len(RESULT_COLUMNS), "4F46E5")

    # Data rows
    for i, r in enumerate(results, 1):
        row_num = i + 2
        ws_results.row_dimensions[row_num].height = 20
        data = [i, r.url, r.vendor, r.product, r.edition, r.version, r.licence_metric, r.eos, r.eol]
        for col_idx, value in enumerate(data, 1):
            ws_results.cell(row=row_num, column=col_idx, value=value)
        _style_data_row(ws_results, row_num, len(RESULT_COLUMNS), alternate=(i % 2 == 0))
        # URL column — blue hyperlink style
        url_cell = ws_results.cell(row=row_num, column=2)
        url_cell.font = Font(color="4F46E5", size=10, name="Calibri", underline="single")

    _auto_fit_columns(ws_results)
    # Freeze header rows
    ws_results.freeze_panes = "A3"

    # -----------------------------------------------------------------------
    # Sheet 2: Errors
    # -----------------------------------------------------------------------
    ws_errors = wb.create_sheet(title="Error URLs")
    ws_errors.sheet_view.showGridLines = False
    ws_errors.row_dimensions[1].height = 20
    ws_errors.row_dimensions[2].height = 30

    # Title
    ws_errors.merge_cells(f"A1:{get_column_letter(len(ERROR_COLUMNS))}1")
    err_title = ws_errors.cell(row=1, column=1)
    err_count = len(errors)
    err_title.value = f"Failed URLs ({err_count} error{'s' if err_count != 1 else ''}) — {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    err_title.font = Font(bold=True, size=13, color="7F1D1D", name="Calibri")
    err_title.fill = PatternFill("solid", fgColor="FEE2E2")
    err_title.alignment = Alignment(horizontal="center", vertical="center")

    # Header
    for col_idx, col_name in enumerate(ERROR_COLUMNS, 1):
        ws_errors.cell(row=2, column=col_idx, value=col_name)
    _style_header_row(ws_errors, 2, len(ERROR_COLUMNS), "DC2626")

    # Data
    for i, e in enumerate(errors, 1):
        row_num = i + 2
        ws_errors.row_dimensions[row_num].height = 20
        data = [i, e.url, e.error]
        for col_idx, value in enumerate(data, 1):
            ws_errors.cell(row=row_num, column=col_idx, value=value)
        _style_data_row(ws_errors, row_num, len(ERROR_COLUMNS), alternate=(i % 2 == 0))
        url_cell = ws_errors.cell(row=row_num, column=2)
        url_cell.font = Font(color="DC2626", size=10, name="Calibri")
        err_cell = ws_errors.cell(row=row_num, column=3)
        err_cell.font = Font(color="7F1D1D", size=10, name="Calibri")

    _auto_fit_columns(ws_errors)
    ws_errors.freeze_panes = "A3"

    # Serialize to bytes
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
